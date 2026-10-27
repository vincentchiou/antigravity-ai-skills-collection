import os
import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

SLUG_MAPPINGS = [
    ("2026海報助理(ChatGPT版)", "poster-assistant-2026-chatgpt", "ChatGPT 2026 電影海報與視覺文宣設計助手"),
    ("Youtube流量封面產生器", "yt-viral-thumbnail-designer", "YouTube 高流量爆款縮圖與視覺封面設計師"),
    ("SEO-YT封面產生器", "seo-yt-thumbnail-designer", "結合搜尋引擎最佳化 (SEO) 與高點擊縮圖設計專家"),
    ("設計你的GPTs/Gems助理", "gpts-gems-prompt-architect", "客製化 GPTs 與 Gemini Gems 系統提示詞架構師"),
    ("格局思維策略顧問", "big-picture-strategic-advisor", "高維度格局思維與頂層戰略決策顧問"),
    ("超好笑梗圖", "hilarious-meme-generator", "網路流行爆笑梗圖與社群迷因生成器"),
    ("YT流量封面產生器", "yt-traffic-thumbnail-generator-gem", "Gemini 版 YouTube 爆款流量縮圖設計專家"),
    ("四格漫畫助理", "four-panel-comic-assistant-gem", "Gemini 版四格漫畫分鏡與幽默故事腳本助理"),
    ("西遊記唐僧爆笑對話(人間清醒毒雞湯)", "journey-to-the-west-tang-monk-toxic-soup", "唐僧反差萌與現代人間清醒毒雞湯對話助手"),
    ("班級好朋友(小朋友用)", "classroom-buddy-for-kids", "專為國小學童設計的溫暖陪伴與班級好朋友"),
    ("SEL孩子輔導專家(教師用)", "sel-child-counseling-expert", "社會情緒學習 (SEL) 與兒少心理輔導專業顧問"),
    ("特教IEP個案產生器", "special-ed-iep-case-generator", "特殊教育個別化教育計畫 (IEP) 個案分析與目標生成器"),
    ("班級學生行為與學習紀錄假資料生成器", "student-behavior-learning-log-mock-data", "擬真班級學生行為與學習評量紀錄假資料生成工具"),
    ("旅遊規劃達人", "travel-itinerary-planner", "客製化深度旅遊行程規劃與在地指南達人"),
    ("動畫分鏡與中文敘事專家 (全能版)", "animation-storyboard-narrative-master", "全能動畫分鏡、鏡頭語言與影視中文敘事專家"),
    ("打造YT影片流量助理", "yt-video-traffic-booster-gem", "Gemini 版 YouTube 影片流量倍增與企劃腳本助理"),
    ("AI Studio TTS 寫稿助手", "ai-studio-tts-scriptwriter-gem", "Gemini 版 Google AI Studio 語音合成 (TTS) 專業講稿寫作助手"),
    ("SRT字幕檔修正工具", "srt-subtitle-corrector-gem", "Gemini 版 SRT 字幕時間軸錯別字與排版修復工具"),
    ("蘇格拉底式教學家教", "socratic-tutoring-mentor", "啟發式提問引導與蘇格拉底式自主學習家教"),
    ("蘇格拉底批判性思考與邏輯驗證助理", "socratic-critical-thinking-verifier", "批判性思維剖析、邏輯謬誤檢驗與論證審查助理"),
    ("蘇格拉底學術期刊審查員", "socratic-academic-peer-reviewer", "高標準學術期刊同儕審查員與論文嚴謹度稽核專家"),
    ("高中歷史遊戲化複習", "high-school-history-gamified-review", "RPG 沉浸式互動與高中歷史遊戲化複習系統"),
    ("2026海報助理(Gem版)", "poster-assistant-2026-gem", "Gemini 版 2026 商業活動與研習海報視覺生成助手"),
    ("新聞稿助理2026(GEM)", "press-release-assistant-2026-gem", "Gemini 版 2026 專業公關新聞稿與媒體發稿助理"),
    ("3.1 TTS -Podcast產生器", "tts-podcast-script-generator", "高品質多聲部對談 Podcast 腳本與 TTS 生成器"),
    ("蘇格拉底式提問助理", "socratic-inquiry-assistant", "蘇格拉底式思維穿透提問與對話啟發助理"),
    ("CRIT 邏輯驗證與反事實思考助理", "crit-logic-counterfactual-thinking", "CRIT 邏輯檢驗、盲點突破與反事實思維演練工具"),
    ("形成性回饋專家", "formative-feedback-expert", "教育評量、課堂形成性回饋與成長型思維導引專家"),
    ("正向思考助理", "positive-mindset-coach", "認知重塑、心理韌性與日常正向賦能教練"),
    ("G3.1-TTS腳本助理", "g31-tts-script-assistant", "Gemini 3.1 擬真語音腳本與口語化表達生成助理"),
    ("金字塔三層次分析", "pyramid-three-level-analysis", "麥肯錫金字塔結構與現象-原因-本質三層次剖析工具"),
    ("孫子哲談", "sun-tzu-philosophy-dialogues", "孫子兵法現代商業決策、職場處世與哲學對話"),
    ("專家主題座談會", "multi-expert-symposium", "多領域專家圓桌思維碰撞與跨界決策座談會"),
    ("牛頓", "isaac-newton-persona", "艾薩克·牛頓爵士 (Sir Isaac Newton) 擬真科學思維化身"),
    ("班級經營神隊友", "classroom-management-wingman", "班級經營、親師溝通、班規建立與突發情境教戰手冊"),
    ("會議記錄助手", "meeting-minutes-summarizer", "結構化會議記錄整理、決議追蹤與行動清單生成器"),
    ("對練學習小幫手", "sparring-learning-companion", "角色扮演情境模擬與自適應對練學習小幫手"),
    ("媒體行銷貼文助手2", "media-marketing-post-assistant", "社群平台高轉換爆款貼文、Hook 標題與文案助手"),
    ("花蓮教育處新聞稿助手2", "hualien-edu-press-release-assistant", "花蓮縣教育處政令宣導、校園成果官方新聞稿專用助手"),
    ("思維學習大師", "master-thinking-learning-coach", "整合元認知、心智模型與終身學習策略思維大師"),
    ("專業繪圖提示詞助手2", "midjourney-dalle-prompt-craftsman", "Midjourney / DALL-E 3 頂級視覺繪圖提示詞工程師"),
    ("老子模擬器", "lao-tzu-taoist-simulator", "老子《道德經》無為而治、道法自然哲學模擬器"),
    ("中西哲人座談會V2", "east-west-philosophers-roundtable-v2", "孔孟老莊與蘇格拉底、康德跨越時空的中西哲人大對談"),
    ("造句擴充小老師", "sentence-expansion-elementary-tutor", "國小生五感觀察修辭、造句擴充與寫作啟蒙小老師"),
    ("九九乘法小老師", "multiplication-table-fun-tutor", "國小九九乘法趣味圖像記憶與口訣遊戲互動小老師"),
    ("依龍‧馬克克", "elon-musk-first-principles-persona", "依龍·馬克克 (Elon Musk) 第一原理思考與第一性思維商業狂人"),
    ("華倫・巴菲特", "warren-buffett-investing-mentor", "華倫·巴菲特 (Warren Buffett) 價值投資、護城河與複利心法導師"),
    ("超級黃金圈Pro", "super-golden-circle-pro", "Simon Sinek 黃金圈法則 (Why-How-What) 深度穿透與價值提煉"),
    ("瓊瑤黃金圈說服大師Pro.V5", "qiongyao-golden-circle-persuasion-v5", "戲劇張力情緒渲染與黃金圈邏輯說服頂級演講文案大師"),
    ("知思行心法", "know-think-act-methodology", "認知-思考-行動三位一體知行合一知思行教練"),
    ("斯普瑞夫 SPREF 知識單元圖譜專家", "spref-knowledge-graph-expert", "SPREF 知識結構拆解、概念單元化與認知圖譜構建專家"),
    ("PBL專案導向學習顧問", "pbl-project-learning-consultant", "PBL 專案式學習 (Project-Based Learning) 課程設計專家"),
    ("國小3年級出題小幫手(請置入學校的電子書)", "grade3-elementary-quiz-generator", "國小三年級各科素養導向評量與試題生成小幫手"),
    ("國中小疑難雜症教育專家", "k9-education-troubleshooter", "國中小學生各類行為、學習障礙與親師疑難排解專家"),
    ("向上管理說話專家", "managing-up-communication-expert", "職場高情商溝通、預期管理與向上報告說話藝術專家"),
    ("與孩子對話", "parent-child-empathic-dialogue", "同理心傾聽、非暴力溝通 (NVC) 與親職深度對話顧問"),
    ("大覺者-理解世間萬物《七維理解模型》", "seven-dimensional-understanding-model", "萬事萬物底層規律《七維理解模型》深度認知解析器"),
    ("四格漫畫助理", "four-panel-comic-assistant-chatgpt", "ChatGPT 版四格漫畫分鏡生成與角色對白企劃助理"),
    ("視覺敘事與電影分鏡高級專家(含生圖動畫音樂指令)", "cinematic-storyboard-visual-narrative-master", "電影級分鏡敘事、Prompt、Runway 動畫與配樂指令全包專家"),
    ("AI Studio TTS 寫稿助手", "ai-studio-tts-scriptwriter-chatgpt", "ChatGPT 版 Google AI Studio TTS 口語劇本寫作助手"),
    ("Youtube字幕檔SRT修正工具", "srt-subtitle-corrector-chatgpt", "ChatGPT 版 YouTube SRT 字幕時間戳與口誤快速清理工具"),
    ("Youtube打造SEO影片流量助理", "yt-seo-video-traffic-booster-chatgpt", "ChatGPT 版 YouTube SEO 關鍵字演算法與高流量文案軍師")
]

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_hyperlink(paragraph, url, text, color="1D4ED8", underline=True):
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w")} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{r_id}" w:history="1"/>')
    new_run = parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:rFonts w:ascii="微軟正黑體" w:hAnsi="微軟正黑體" w:eastAsia="微軟正黑體"/><w:color w:val="{color}"/>{"<w:u w:val=\'single\'/>" if underline else ""}</w:rPr><w:t>{text}</w:t></w:r>')
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)

def create_document():
    doc = Document()
    
    # Page Setup (A4, standard margins)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = '微軟正黑體'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')

    # Color Palette
    PRIMARY = RGBColor(0x0F, 0x17, 0x2A)   # Dark Slate 900
    SECONDARY = RGBColor(0x1E, 0x3A, 0x8A) # Blue 900
    ACCENT = RGBColor(0x25, 0x63, 0xEB)    # Blue 600
    TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)# Slate 500

    # Title Section
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("🤖 AI 助理與專家技能總庫手冊")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = PRIMARY

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(14)
    sub_run = subtitle_p.add_run("涵蓋 62 款 Gemini Gems & ChatGPT GPTs 之 Agent Skills 轉換清單、調用指南與實戰手冊")
    sub_run.font.size = Pt(12.5)
    sub_run.font.color.rgb = SECONDARY

    # Meta box
    meta_tbl = doc.add_table(rows=1, cols=1)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = meta_tbl.cell(0, 0)
    set_cell_background(c, "F1F5F9")
    set_cell_margins(c, top=140, bottom=140, left=200, right=200)
    c_p = c.paragraphs[0]
    c_p.paragraph_format.space_after = Pt(0)
    c_p.paragraph_format.line_spacing = 1.3
    m_run = c_p.add_run("📌 本地技能庫目錄：h:\\我的雲端硬碟\\Agent\\project\\AI助理\\\n"
                      "📌 原始資料來源：🤖 助理機器人 Prompt 資源庫 (Notion Database)\n"
                      "📌 技能標準規範：Antigravity / Claude Code YAML Frontmatter + SKILL.md 架構\n"
                      "📌 技能收錄總數：62 款（含教學班級、視覺繪圖、影音TTS、高階哲學思維、職場行政）")
    m_run.font.size = Pt(9.5)
    m_run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 1: How to Invoke
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r1 = h1.add_run("一、 技能系統架構與如何呼叫 (How to Invoke Skills)")
    r1.font.size = Pt(15)
    r1.font.bold = True
    r1.font.color.rgb = PRIMARY

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    p.add_run("所有轉化後的助理皆採用現代 AI Agent 標準的 ")
    r_code = p.add_run("SKILL.md")
    r_code.font.bold = True
    r_code.font.color.rgb = ACCENT
    p.add_run(" 封裝規範。每個技能資料夾皆包含 YAML Frontmatter（內含 name 與 description）以及完整的 System Prompt 核心指令。在 Antigravity 與各類智能體環境中，您可以透過以下三種方式呼叫：")

    # Call method 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("1. 自然語意自動觸發（最直覺、最推薦）")
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = SECONDARY

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(4)
    p.add_run("您無須死記英文代稱，直接在日常對話中提出需求，Agent 的漸進式認知系統（Progressive Disclosure）會依據情境自動比對並掛載對應技能：\n"
              "• 「我下週要辦一場國小 AI 研習，請用 ")
    p.add_run("2026海報助理").font.bold = True
    p.add_run(" 幫我產出 Midjourney 生圖提示詞與宣傳大綱。」\n"
              "• 「班上有個孩子情緒起伏很大，請以 ")
    p.add_run("SEL孩子輔導專家").font.bold = True
    p.add_run(" 的角度給我具體的同理回應與引導四步驟。」\n"
              "• 「請用 ")
    p.add_run("金字塔三層次分析").font.bold = True
    p.add_run(" 幫我深入剖析為什麼現代人普遍產生短影音成癮現象。」")

    # Call method 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("2. 明確指定 Skill Slug 代稱（精確命中）")
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = SECONDARY

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(4)
    p.add_run("當您想確保 100% 精準執行特定助理時，可直接在對話中提及該 Skill 的 Slug 代稱：\n"
              "• 「請啟動 ")
    p.add_run("`socratic-critical-thinking-verifier`").font.bold = True
    p.add_run("，檢驗以下教育政見報告的推論有效性與潛在邏輯謬誤。」\n"
              "• 「請調用 ")
    p.add_run("`yt-viral-thumbnail-designer`").font.bold = True
    p.add_run("，為這支教學影片設計 3 款高點擊率封面構圖方案。」")

    # Call method 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("3. 指定讀取 SKILL.md 檔案路徑（強制照章執行）")
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = SECONDARY

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(8)
    p.add_run("在對話中直接要求 Agent 閱讀特定檔案，例如：「請直接參閱 ")
    p.add_run("./special-ed-iep-case-generator/SKILL.md").font.bold = True
    p.add_run("，並嚴格按照其中的評量指標為注意力缺失學生擬定新學期個別化教育計畫。」")

    # Global skill tip box
    tip_tbl = doc.add_table(rows=1, cols=1)
    tip_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tip_tbl.cell(0, 0)
    set_cell_background(c, "EFF6FF")
    set_cell_margins(c, top=140, bottom=140, left=200, right=200)
    cp = c.paragraphs[0]
    cp.paragraph_format.space_after = Pt(0)
    r_tip_title = cp.add_run("💡 如何設為「全域技能」（在任何專案或任何路徑隨時隨地可用）：\n")
    r_tip_title.font.bold = True
    r_tip_title.font.size = Pt(10)
    r_tip_title.font.color.rgb = SECONDARY
    r_tip_body = cp.add_run(
        "目前這 62 款技能存放於本專案資料夾。若您希望在電腦上的任何專案視窗中皆能隨時調用，"
        "只需將任一技能資料夾複製或建立軟連結（Symlink）至 Antigravity 全域技能目錄：\n"
        "👉 全域路徑：C:\\Users\\user\\.gemini\\config\\skills\\\n"
        "（您現有的 visual-infographic-deck 與 bambu-3d-print-designer 亦存放於該目錄中）。"
    )
    r_tip_body.font.size = Pt(9.5)
    r_tip_body.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 2: Category Overview
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r2 = h2.add_run("二、 技能五大領域分類與統計分析")
    r2.font.size = Pt(15)
    r2.font.bold = True
    r2.font.color.rgb = PRIMARY

    cat_tbl = doc.add_table(rows=6, cols=3)
    cat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cat_headers = ["領域分類", "收錄數量", "核心涵蓋功能與適用場景"]
    for i, h in enumerate(cat_headers):
        cell = cat_tbl.cell(0, i)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10)

    cat_data = [
        ("🎓 教學與班級經營", "18 款", "涵蓋親師溝通、班級經營神隊友、SEL情緒輔導、特教IEP評量、國小三年級出題、假資料生成、造句擴充、九九乘法及歷史遊戲化複習。"),
        ("🎨 視覺設計與多媒體", "14 款", "2026電影海報(Gemini/ChatGPT)、YouTube爆款縮圖封面、四格漫畫助理、動畫分鏡語言、頂級Midjourney/DALL-E提示詞設計。"),
        ("🎙️ 語音合成與字幕", "6 款", "Google AI Studio TTS專業講稿助手、Podcast雙人對話腳本、SRT字幕時間戳校對與錯別字批次修正工具。"),
        ("🧠 高階思維與哲人對談", "14 款", "蘇格拉底式提問與批判思維、金字塔三層次剖析、超級黃金圈Pro、瓊瑤說服大師、老子模擬器、孫子哲談、七維理解模型及牛頓擬真化身。"),
        ("💼 行政、企劃與公關", "10 款", "花蓮教育處官方新聞稿助手、公關發稿助理2026、結構化會議記錄、爆款社群貼文、向上管理說話藝術、旅遊規劃達人。")
    ]

    for row_idx, data in enumerate(cat_data, 1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = cat_tbl.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            if col_idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = SECONDARY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3: Master Catalog Table
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r3 = h3.add_run("三、 全 62 款 AI 助理與專家技能完整對照清單")
    r3.font.size = Pt(15)
    r3.font.bold = True
    r3.font.color.rgb = PRIMARY

    backup_path = os.path.join('_source_data', 'notion_assistants_backup.json')
    with open(backup_path, 'r', encoding='utf-8') as f:
        extracted = json.load(f)
    valid_pages = [p for p in extracted if p.get('名稱')]

    tbl = doc.add_table(rows=len(SLUG_MAPPINGS) + 1, cols=6)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["序號", "技能名稱 (Title)", "Skill Slug (目錄名)", "平台來源", "核心定位與功能描述", "原始連結"]
    for i, h in enumerate(headers):
        cell = tbl.cell(0, i)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, top=140, bottom=140, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9)

    for i in range(len(SLUG_MAPPINGS)):
        raw_title, slug, desc = SLUG_MAPPINGS[i]
        page = valid_pages[i]
        url = page.get('參考網址', '').strip()
        tags_str = page.get('標籤', '').strip()
        
        if 'gemini.google.com' in url or 'Gem' in tags_str or '(Gem版)' in raw_title or '(GEM)' in raw_title:
            platform = "Gemini Gem"
        else:
            platform = "ChatGPT GPT"

        row_idx = i + 1
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"

        # Col 0: Index
        c0 = tbl.cell(row_idx, 0)
        set_cell_background(c0, bg)
        set_cell_margins(c0, top=70, bottom=70, left=50, right=50)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p0.add_run(f"{row_idx:02d}")
        r.font.size = Pt(8.5)
        r.font.color.rgb = TEXT_MUTED

        # Col 1: Title
        c1 = tbl.cell(row_idx, 1)
        set_cell_background(c1, bg)
        set_cell_margins(c1, top=70, bottom=70, left=60, right=60)
        p1 = c1.paragraphs[0]
        r = p1.add_run(raw_title)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = PRIMARY

        # Col 2: Slug
        c2 = tbl.cell(row_idx, 2)
        set_cell_background(c2, bg)
        set_cell_margins(c2, top=70, bottom=70, left=60, right=60)
        p2 = c2.paragraphs[0]
        r = p2.add_run(slug)
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        r.font.color.rgb = ACCENT

        # Col 3: Platform
        c3 = tbl.cell(row_idx, 3)
        set_cell_background(c3, bg)
        set_cell_margins(c3, top=70, bottom=70, left=50, right=50)
        p3 = c3.paragraphs[0]
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p3.add_run(platform)
        r.font.size = Pt(8)
        if "Gemini" in platform:
            r.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)
        else:
            r.font.color.rgb = RGBColor(0x7C, 0x3A, 0xED)

        # Col 4: Desc
        c4 = tbl.cell(row_idx, 4)
        set_cell_background(c4, bg)
        set_cell_margins(c4, top=70, bottom=70, left=60, right=60)
        p4 = c4.paragraphs[0]
        r = p4.add_run(desc)
        r.font.size = Pt(8)

        # Col 5: URL
        c5 = tbl.cell(row_idx, 5)
        set_cell_background(c5, bg)
        set_cell_margins(c5, top=70, bottom=70, left=50, right=50)
        p5 = c5.paragraphs[0]
        p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if url:
            add_hyperlink(p5, url, "前往", color="2563EB", underline=True)
        else:
            r = p5.add_run("-")
            r.font.size = Pt(8)
            r.font.color.rgb = TEXT_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 4: Showcase & Methodology
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r4 = h4.add_run("四、 代表性技能深度拆解與提示詞實戰案例")
    r4.font.size = Pt(15)
    r4.font.bold = True
    r4.font.color.rgb = PRIMARY

    showcase_items = [
        ("1. 設計你的GPTs/Gems助理 (gpts-gems-prompt-architect)",
         "專門協助使用者打造專屬 AI 助理的元架構師 (Meta-Prompt Architect)。"
         "採用五階段引導模式：釐清目標受眾 -> 互動層級 -> 步驟式執行SOP -> 防呆與邊界約束 -> 開場白與知識庫規劃。可快速生成企業或教學現場專用的高水準 System Instructions。"),
        ("2. 2026海報助理 (poster-assistant-2026-chatgpt / poster-assistant-2026-gem)",
         "突破傳統文字文宣限制，定位為「AI 海報視覺總監＋生圖 Prompt 設計師」。"
         "輸出直接支援 SD / Midjourney / ComfyUI / Z-Image 的專業 Prompt 結構：包含【主畫面】（場景＋角色＋情緒＋光影）、【風格與光影細節】、【行銷文案排版配置】。"),
        ("3. SEL孩子輔導專家(教師用) (sel-child-counseling-expert)",
         "針對班級教學現場情緒與人際衝突設計。遵循社會情緒學習 (SEL) 五大核心能力，"
         "引導老師依序落實：同理接納幼兒情緒 -> 澄清行為背後的未滿足需求 -> 引導覺察身體感受 -> 共同討論正向解方，避免說教式語言造成學生防衛心理。"),
        ("4. 蘇格拉底批判性思考與邏輯驗證助理 (socratic-critical-thinking-verifier)",
         "以啟發式詰問與嚴密邏輯檢驗為核心。主動辨識論述中的假前提、以偏概全、滑坡謬誤或虛假二分法，"
         "透過多輪深層提問，迫使對話者反思自我思考盲區，培育獨立思考與高階論證能力。")
    ]

    for title_text, desc_text in showcase_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title_text)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = SECONDARY

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.line_spacing = 1.2
        p_desc.paragraph_format.space_after = Pt(6)
        r = p_desc.add_run(desc_text)
        r.font.size = Pt(9.5)

    # Save document
    out_path = r"h:\我的雲端硬碟\Agent\project\AI助理\AI助理與專家技能總庫手冊.docx"
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_document()
